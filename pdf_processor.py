import os
import fitz  # PyMuPDF
from docx import Document


# ============================================================
# CASETAG - PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(pdf_path):

    if not os.path.exists(pdf_path):
        raise FileNotFoundError(
            f"PDF not found:\n{pdf_path}"
        )

    print("\nOpening PDF...")

    document = fitz.open(pdf_path)

    print(f"Total pages: {len(document)}")

    pages = []

    for page_number, page in enumerate(document, start=1):

        text = page.get_text("text")

        if text:
            text = text.strip()

        if text:
            pages.append(
                f"\n--- PAGE {page_number} ---\n{text}"
            )

    document.close()

    extracted_text = "\n".join(pages)

    return extracted_text


# ============================================================
# CASETAG - WORD DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(docx_path):

    if not os.path.exists(docx_path):
        raise FileNotFoundError(
            f"Word document not found:\n{docx_path}"
        )

    print("\nOpening Word document...")

    document = Document(docx_path)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # --------------------------------------------------------
    # Extract text from tables also
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                paragraphs.append(" | ".join(row_text))

    extracted_text = "\n".join(paragraphs)

    return extracted_text


# ============================================================
# CASETAG - OLD WORD DOC TEXT EXTRACTION
# ============================================================

def extract_text_from_doc(doc_path):

    if not os.path.exists(doc_path):
        raise FileNotFoundError(
            f"Word document not found:\n{doc_path}"
        )

    raise ValueError(
        "Old .doc files are not supported directly. "
        "Please save the Word document as .docx and upload it again."
    )


# ============================================================
# CASETAG - TXT TEXT EXTRACTION
# ============================================================

def extract_text_from_txt(txt_path):

    if not os.path.exists(txt_path):
        raise FileNotFoundError(
            f"TXT file not found:\n{txt_path}"
        )

    print("\nOpening TXT file...")

    # --------------------------------------------------------
    # First try UTF-8
    # --------------------------------------------------------

    try:

        with open(
            txt_path,
            "r",
            encoding="utf-8"
        ) as file:

            extracted_text = file.read()

    except UnicodeDecodeError:

        # ----------------------------------------------------
        # Fallback encoding
        # ----------------------------------------------------

        with open(
            txt_path,
            "r",
            encoding="latin-1"
        ) as file:

            extracted_text = file.read()

    return extracted_text


# ============================================================
# CASETAG - GENERIC DOCUMENT TEXT EXTRACTION
# ============================================================

def extract_text_from_document(file_path):

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found:\n{file_path}"
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()

    print(
        f"\nDetected file type: {extension}"
    )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    elif extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    # --------------------------------------------------------
    # DOC
    # --------------------------------------------------------

    elif extension == ".doc":

        return extract_text_from_doc(
            file_path
        )

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    elif extension == ".txt":

        return extract_text_from_txt(
            file_path
        )

    # --------------------------------------------------------
    # Unsupported
    # --------------------------------------------------------

    else:

        raise ValueError(
            "Unsupported file format.\n"
            "Please upload PDF, DOCX, or TXT."
        )


# ============================================================
# CLEAN EXTRACTED TEXT
# ============================================================

def clean_text(text):

    lines = text.splitlines()

    cleaned_lines = []

    for line in lines:

        line = line.strip()

        if not line:
            continue

        cleaned_lines.append(line)

    return "\n".join(cleaned_lines)


# ============================================================
# SAVE EXTRACTED TEXT
# ============================================================

def save_extracted_text(
    text,
    output_path="data/extracted_case.txt"
):

    directory = os.path.dirname(
        output_path
    )

    if directory:
        os.makedirs(
            directory,
            exist_ok=True
        )

    with open(
        output_path,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(text)

    print(
        f"\nExtracted text saved to:\n{output_path}"
    )


# ============================================================
# DISPLAY PREVIEW
# ============================================================

def display_preview(text):

    print(
        "\n"
        + "=" * 70
    )

    print(
        "              EXTRACTED CASE TEXT"
    )

    print(
        "=" * 70
    )

    print(
        f"\nTotal characters: {len(text)}"
    )

    print(
        f"Total words: {len(text.split())}"
    )

    print(
        "\nPreview:\n"
    )

    preview = text[:3000]

    print(preview)

    if len(text) > 3000:

        print(
            "\n... [preview truncated]"
        )


# ============================================================
# MAIN
# ============================================================

def main():

    print(
        "=" * 70
    )

    print(
        "             CASETAG DOCUMENT PROCESSOR"
    )

    print(
        "=" * 70
    )

    file_path = input(
        "\nEnter the path of your case document "
        "(PDF / DOC / DOCX / TXT):\n"
    ).strip().strip('"')

    if not file_path:

        print(
            "\nNo file path provided."
        )

        return

    try:

        # ----------------------------------------------------
        # Extract
        # ----------------------------------------------------

        extracted_text = extract_text_from_document(
            file_path
        )

        if not extracted_text.strip():

            print(
                "\nNo text could be extracted from this document."
            )

            print(
                "The file may be empty or image-based."
            )

            return

        # ----------------------------------------------------
        # Clean
        # ----------------------------------------------------

        cleaned_text = clean_text(
            extracted_text
        )

        # ----------------------------------------------------
        # Preview
        # ----------------------------------------------------

        display_preview(
            cleaned_text
        )

        # ----------------------------------------------------
        # Save
        # ----------------------------------------------------

        save_extracted_text(
            cleaned_text
        )

        print(
            "\n"
            + "=" * 70
        )

        print(
            "          DOCUMENT PROCESSING COMPLETE"
        )

        print(
            "=" * 70
        )

    except Exception as e:

        print(
            "\nERROR:"
        )

        print(e)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()